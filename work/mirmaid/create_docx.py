from docx import Document
from docx.shared import Inches
from docx.enum.text import WD_ALIGN_PARAGRAPH

# Create document
doc = Document()

# Title
title = doc.add_heading("Machine Learning", level=1)
title.alignment = WD_ALIGN_PARAGRAPH.CENTER

# Text
doc.add_heading("Introduction", level=2)

doc.add_paragraph(
    "Machine learning is a branch of artificial intelligence "
    "that enables computers to learn from data and make predictions."
)

# Diagram heading
doc.add_heading("Machine Learning Process", level=2)

doc.add_paragraph(
    "The following diagram represents the basic machine learning workflow:"
)

# Insert Mermaid-generated PNG
image = doc.add_picture(
    "diagram.png",
    width=Inches(3.1)
)

# Center the image
last_paragraph = doc.paragraphs[-1]
last_paragraph.alignment = WD_ALIGN_PARAGRAPH.CENTER

# Explanation after diagram
doc.add_paragraph(
    "The process starts with input data, followed by preprocessing "
    "and model training. The trained model is then used to generate predictions."
)

# Save document
doc.save("machine_learnin_notes.docx")

print("DOCX generated successfully!")