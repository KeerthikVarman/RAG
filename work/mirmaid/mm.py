from docx import Document
from docx.shared import Inches

doc = Document()

doc.add_heading("Machine Learning Process", level=1)

doc.add_paragraph(
    "The following diagram represents the machine learning workflow:"
)

doc.add_picture(
    "diagram.png",
    width=Inches(6)
)

doc.save("machine_learnin_notes.docx")